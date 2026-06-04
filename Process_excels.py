#   --force  : recopie les templates même si un Excel existe déjà
#   --ts     : ajoute un timestamp au nom des Excel créés

import os, sys, csv, glob, shutil, io, argparse, json
from pathlib import Path
from datetime import datetime

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    print("openpyxl manquant. Installe: python -m pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent

def sniff_delimiter(sample: str, default=","):
    candidates = [";", ",", "\t", "|"]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        return dialect.delimiter
    except Exception:
        counts = {d: sample.count(d) for d in candidates}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else default

def read_csv_file(path: Path):
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = "\n".join(text.splitlines()[:30])
    delim = sniff_delimiter(sample)
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    return list(reader)

def headers_row(ws, start_row=1):
    headers = {}
    if ws.max_row >= start_row:
        for cell in ws[start_row]:
            if cell.value:
                headers[str(cell.value).strip()] = cell.col_idx
    return headers

def write_headers(ws, header_names, start_row=1):
    for idx, h in enumerate(header_names, start=1):
        ws.cell(row=start_row, column=idx, value=h)

def append_rows(ws, rows, header_map=None, start_row=1):
    header_map = header_map or {}
    excel_headers = headers_row(ws, start_row=start_row)

    # Si aucune ligne d'entêtes détectée à start_row, on la crée à partir du 1er CSV
    if not excel_headers:
        first = rows[0] if rows else {}
        header_names = [header_map.get(k, k) for k in first.keys()]
        if header_names:
            write_headers(ws, header_names, start_row=start_row)
            excel_headers = headers_row(ws, start_row=start_row)

    for r in rows:
        out = [None] * max(1, len(excel_headers))
        for csv_key, value in r.items():
            excel_key = header_map.get(csv_key, csv_key)
            if excel_key in excel_headers:
                out[excel_headers[excel_key] - 1] = value
        # Ajoute en bas du tableau (première ligne vide après ce qui existe)
        ws.append(out)

def ensure_table(ws, table_name, start_row=1):
    # Pose une table seulement s'il y a au moins entêtes + 1 ligne
    if ws.tables:
        return
    max_row = ws.max_row
    max_col = ws.max_column
    if max_row < start_row + 1 or max_col < 1:
        return
    ref = f"A{start_row}:{get_column_letter(max_col)}{max_row}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
    ws.add_table(t)

def copy_template_if_needed(src: Path, dst: Path, force=False, timestamp=False):
    dst.parent.mkdir(parents=True, exist_ok=True)
    target = dst
    if timestamp:
        target = dst.with_name(dst.stem + "_" + datetime.now().strftime("%Y%m%d_%H%M%S") + dst.suffix)
    # if --force => recopier toujours; sinon, utiliser la plus récente si déjà présente
    pattern = dst.parent / f"{dst.stem}*{dst.suffix}"
    existing = sorted(pattern.parent.glob(pattern.name), key=lambda p: p.stat().st_mtime, reverse=True)
    if force or not existing:
        shutil.copy2(src, target)
        print(f"[COPY] {src.name} → {target.relative_to(ROOT)}")
        return target
    else:
        print(f"[KEEP] Utilisation du classeur existant : {existing[0].relative_to(ROOT)}")
        return existing[0]

def load_rows_from_globs(patterns: str):
    rows, files = [], []
    parts = [p.strip() for p in patterns.split(";") if p.strip()]
    for p in parts:
        files.extend(glob.glob(str(ROOT / p)))
    for f in files:
        try:
            rows.extend(read_csv_file(Path(f)))
        except Exception as e:
            print(f"[WARN] Lecture CSV échouée: {f} → {e}")
    return rows, files

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="Recopier les templates même si un Excel existe déjà")
    ap.add_argument("--ts", action="store_true", help="Ajouter un timestamp aux fichiers Excel créés")
    ap.add_argument("--config", default="modules_config.json", help="Chemin du fichier JSON de config")
    args = ap.parse_args()

    cfg_path = ROOT / args.config
    if not cfg_path.exists():
        print(f"Config introuvable: {cfg_path}")
        sys.exit(1)
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    modules = cfg.get("modules", [])

    for mod in modules:
        name        = mod.get("name", "module")
        template    = ROOT / mod["template"]
        # on conserve 'out_file' s'il est présent, mais on le remplacera par le nom du CSV
        out_file    = ROOT / mod.get("out_file", "Excel/Output.xlsx")
        sheet       = mod.get("sheet", "Data")
        csv_glob    = mod.get("csv_glob", "")
        start_row   = int(mod.get("start_row", 1))
        header_map  = mod.get("header_map", {})

        if not template.exists():
            print(f"[{name}] Template introuvable : {template.relative_to(ROOT)}")
            continue

        # 1) Charger les CSV (et lister les fichiers)
        data_all, files = load_rows_from_globs(csv_glob)
        if not files:
            print(f"[{name}] Aucun CSV trouvé pour: {csv_glob}")
            continue
        print(f"[{name}] {len(files)} fichier(s) CSV trouvés.")

        # 2) Si plusieurs CSV → un Excel par CSV ; si un seul → nommer l'Excel comme le CSV
        if len(files) == 1:
            csv_path = Path(files[0])
            csv_name = csv_path.stem
            out_file = (ROOT / "Excel" / f"{csv_name}.xlsx")
            out_path = copy_template_if_needed(template, out_file, force=args.force, timestamp=args.ts)

            # relire UNIQUEMENT ce CSV (évite l'agrégat de load_rows_from_globs)
            try:
                data = read_csv_file(csv_path)
            except Exception as e:
                print(f"[{name}] WARN: lecture CSV échouée: {csv_path.name} → {e}")
                continue

            try:
                wb = load_workbook(str(out_path))
            except Exception as e:
                print(f"[{name}] ERREUR: échec ouverture {out_path.name} → {e}")
                continue

            ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(title=sheet)

            if data:
                append_rows(ws, data, header_map=header_map, start_row=start_row)
                try:
                    ensure_table(ws, table_name=f"Table_{sheet.replace(' ','_')[:20]}", start_row=start_row)
                except Exception as e:
                    print(f"[{name}] WARN: table non posée ({e})")

            try:
                wb.save(str(out_path))
                print(f"[{name}] ✓ {csv_path.name} → {out_path.name} / Feuille: {sheet}")
            except Exception as e:
                print(f"[{name}] ERREUR: sauvegarde {out_path.name} → {e}")

        else:
            # Plusieurs CSV : traiter chaque fichier séparément
            for f in files:
                csv_path = Path(f)
                csv_name = csv_path.stem
                out_file = (ROOT / "Excel" / f"{csv_name}.xlsx")
                out_path = copy_template_if_needed(template, out_file, force=args.force, timestamp=args.ts)

                try:
                    data = read_csv_file(csv_path)
                except Exception as e:
                    print(f"[{name}] WARN: lecture CSV échouée: {csv_path.name} → {e}")
                    continue

                try:
                    wb = load_workbook(str(out_path))
                except Exception as e:
                    print(f"[{name}] ERREUR: échec ouverture {out_path.name} → {e}")
                    continue

                ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(title=sheet)

                if data:
                    append_rows(ws, data, header_map=header_map, start_row=start_row)
                    try:
                        ensure_table(ws, table_name=f"Table_{sheet.replace(' ','_')[:20]}", start_row=start_row)
                    except Exception as e:
                        print(f"[{name}] WARN: table non posée ({e})")

                try:
                    wb.save(str(out_path))
                    print(f"[{name}] ✓ {csv_path.name} → {out_path.name} / Feuille: {sheet}")
                except Exception as e:
                    print(f"[{name}] ERREUR: sauvegarde {out_path.name} → {e}")

    print("\nTerminé. Les classeurs sont dans le dossier Excel/.")

if __name__ == "__main__":
    main()
